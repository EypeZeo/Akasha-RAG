"""
PR2C-1 回归测试：批量导出的链接字段必须来自 canonical_url，不能硬编码 douyin.com

BUG-05：全部 6 条批量导出路径此前都硬编码
`https://www.douyin.com/video/{platform_item_id}`——B 站内容导出出来的
链接必然指向错误的域名。修复后统一走 `display_link(fv)`。这里对每条
格式各验证一次（不只测 Markdown 和 Word）；PDF 路径的链接正确性在
`test_batch_export_pdf_escaping.py` 里和转义一起验证。
"""
from __future__ import annotations

import io
from types import SimpleNamespace

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.services.batch_export_service import BatchExportService


def _fake_item(platform_item_id: str, canonical_url: str, author: str = "某UP主", title: str | None = None):
    cache = SimpleNamespace(
        platform_item_id=platform_item_id,
        title=title or f"标题-{platform_item_id}",
        updated_at=None,
        transcript_text="这是转写正文。",
    )
    fv = SimpleNamespace(author=author, canonical_url=canonical_url)
    return cache, fv


def test_markdown_export_uses_canonical_url_not_douyin():
    service = BatchExportService()
    item = _fake_item("BV1xx", "https://www.bilibili.com/video/BV1xx")
    md = service._render_item_markdown(*item, content_type="original")
    assert "bilibili.com" in md
    assert "douyin.com" not in md


def test_markdown_export_shows_placeholder_when_link_missing():
    service = BatchExportService()
    item = _fake_item("BV1xx", "")
    md = service._render_item_markdown(*item, content_type="original")
    assert "（链接缺失）" in md
    assert "douyin.com" not in md


def test_excel_export_uses_canonical_url_not_douyin():
    service = BatchExportService()
    buf = service._export_excel([_fake_item("BV1xx", "https://www.bilibili.com/video/BV1xx")], content_type="original")
    wb = load_workbook(buf)
    ws = wb.active
    row_values = [str(ws.cell(row=2, column=col).value or "") for col in range(1, ws.max_column + 1)]
    joined = " ".join(row_values)
    assert "bilibili.com" in joined
    assert "douyin.com" not in joined


def test_word_single_export_uses_canonical_url_not_douyin():
    service = BatchExportService()
    buf = service._export_word_single([_fake_item("BV1xx", "https://www.bilibili.com/video/BV1xx")], content_type="original")
    doc = Document(io.BytesIO(buf.getvalue()))
    table_text = " ".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "bilibili.com" in table_text
    assert "douyin.com" not in table_text


def test_word_zip_export_uses_canonical_url_not_douyin():
    service = BatchExportService()
    zip_buf = service._export_word_zip([_fake_item("BV1xx", "https://www.bilibili.com/video/BV1xx")], content_type="original")
    import zipfile
    with zipfile.ZipFile(zip_buf) as zf:
        names = zf.namelist()
        assert len(names) == 1
        docx_bytes = zf.read(names[0])
    doc = Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "bilibili.com" in full_text
    assert "douyin.com" not in full_text


def test_ppt_export_uses_canonical_url_not_douyin():
    service = BatchExportService()
    buf = service._export_ppt([_fake_item("BV1xx", "https://www.bilibili.com/video/BV1xx")], content_type="original")
    prs = Presentation(io.BytesIO(buf.getvalue()))
    all_text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                all_text.append(shape.text_frame.text)
    joined = "\n".join(all_text)
    assert "bilibili.com" in joined
    assert "douyin.com" not in joined
