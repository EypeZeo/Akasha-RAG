"""
BUG-15 回归测试：Excel 导出不能静默截断超长内容

Excel 单元格有约 32767 字符的硬上限。之前的实现直接把完整转写文本塞进
一个单元格，长视频会被 Excel 打开时静默截断/判定为损坏。修复后超长内容
按阈值拆成同一列的多个"延续行"，不引入新列、不依赖打包模式，保证全部
内容都留在这一份 xlsx 里。
"""
from __future__ import annotations

from types import SimpleNamespace

from openpyxl import load_workbook

from app.services.batch_export_service import BatchExportService


def _fake_item(platform_item_id: str, transcript_text: str, author: str = "某作者"):
    cache = SimpleNamespace(
        platform_item_id=platform_item_id,
        title=f"标题-{platform_item_id}",
        updated_at=None,
        transcript_text=transcript_text,
    )
    fv = SimpleNamespace(author=author)
    return cache, fv


def test_oversized_transcript_is_split_across_continuation_rows_without_data_loss():
    service = BatchExportService()
    long_text = "".join(f"第{i}段内容。" for i in range(10000))  # far beyond one Excel cell
    assert len(long_text) > 32767

    buf = service._export_excel([_fake_item("v1", long_text)], content_type="original")

    wb = load_workbook(buf)
    ws = wb.active

    transcript_col = None
    for col in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=col).value == "原始转写文本":
            transcript_col = col
            break
    assert transcript_col is not None

    collected = []
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row=row, column=transcript_col).value
        assert value is None or len(value) <= 30000
        if value:
            collected.append(value)

    assert "".join(collected) == long_text
    assert ws.max_row > 2  # must have produced continuation rows, not one giant cell


def test_normal_length_transcript_stays_on_a_single_row():
    service = BatchExportService()
    short_text = "一段正常长度的转写文本。"

    buf = service._export_excel([_fake_item("v2", short_text)], content_type="original")

    wb = load_workbook(buf)
    ws = wb.active
    # header row + exactly one data row, no continuation rows produced.
    assert ws.max_row == 2
